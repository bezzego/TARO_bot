"""Выгрузка записей в Excel: лист со всеми записями и итоги по месяцам."""
import io
from collections import defaultdict
from datetime import datetime

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import config
import database

STATUS_RU = {
    config.STATUS_WAITING_PAYMENT: "Бронь без чека",
    config.STATUS_CHECKING: "Ждёт проверки оплаты",
    config.STATUS_CONFIRMED: "Подтверждена",
    config.STATUS_REJECTED: "Отклонена",
    config.STATUS_CANCELLED: "Отменена",
}
STATUS_FILL = {
    config.STATUS_CONFIRMED: "D9F2D0",
    config.STATUS_CHECKING: "FFF2B3",
    config.STATUS_WAITING_PAYMENT: "FCE4D6",
    config.STATUS_REJECTED: "EDEDED",
    config.STATUS_CANCELLED: "EDEDED",
}
HEADER_FILL = PatternFill("solid", fgColor="4B2E83")
HEADER_FONT = Font(bold=True, color="FFFFFF")
MONTHS = ["Январь", "Февраль", "Март", "Апрель", "Май", "Июнь", "Июль",
          "Август", "Сентябрь", "Октябрь", "Ноябрь", "Декабрь"]


def _text(value, limit=32000):
    """Текст для ячейки без управляющих символов (с ними xlsx не откроется)."""
    if value is None:
        return ""
    return ILLEGAL_CHARACTERS_RE.sub("", str(value))[:limit]


def _header(ws, titles, widths):
    ws.append(titles)
    for i, w in enumerate(widths, 1):
        c = ws.cell(row=1, column=i)
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"


async def build_bookings_xlsx() -> bytes:
    cur = await database.db.execute(
        "SELECT b.id, b.kind, b.call_minutes, b.status, b.amount, b.num_questions, b.questions, "
        "b.participants, b.story, b.payment_info, "
        "COALESCE(s.date, b.slot_date_cache) AS date, COALESCE(s.time, b.slot_time_cache) AS time, "
        "u.name, u.username, u.phone "
        "FROM bookings b JOIN users u ON u.user_id = b.user_id LEFT JOIN slots s ON s.id = b.slot_id "
        "ORDER BY COALESCE(s.date, b.slot_date_cache) DESC, COALESCE(s.time, b.slot_time_cache) DESC, b.id DESC")
    rows = await cur.fetchall()

    wb = Workbook()
    ws = wb.active
    ws.title = "Записи"
    _header(ws,
            ["№", "Тип", "Длительность, мин", "Дата", "Время", "Клиент", "Username", "Телефон",
             "Статус", "Сумма, ₽", "Вопросов", "Вопросы", "Участники", "История", "Данные перевода"],
            [7, 11, 11, 11, 13, 20, 18, 16, 20, 10, 9, 45, 25, 60, 30])

    monthly = defaultdict(lambda: {"spread": [0, 0], "call": [0, 0]})
    for r in rows:
        is_call = r["kind"] == "call"
        try:
            d = datetime.strptime(r["date"], "%Y-%m-%d").date() if r["date"] else None
        except ValueError:
            d = None
        t = r["time"] or ""
        if is_call and t and r["call_minutes"]:
            t = f"{t}–{database._add_minutes(t, r['call_minutes'])}"
        ws.append([
            r["id"], "Созвон" if is_call else "Расклад", r["call_minutes"] if is_call else None,
            d, t, _text(r["name"]), _text("@" + r["username"] if r["username"] else ""), _text(r["phone"]),
            STATUS_RU.get(r["status"], _text(r["status"])), r["amount"], r["num_questions"],
            _text(r["questions"]), _text(r["participants"]), _text(r["story"]), _text(r["payment_info"]),
        ])
        row = ws.max_row
        for c in ws[row]:
            # Текст клиента, начинающийся с «=», openpyxl записал бы как формулу
            if isinstance(c.value, str) and c.value.startswith("="):
                c.data_type = "s"
        ws.cell(row=row, column=4).number_format = "DD.MM.YYYY"
        fill = STATUS_FILL.get(r["status"])
        if fill:
            ws.cell(row=row, column=9).fill = PatternFill("solid", fgColor=fill)
        for col in (12, 13, 14, 15):
            ws.cell(row=row, column=col).alignment = Alignment(wrap_text=True, vertical="top")
        if r["status"] == config.STATUS_CONFIRMED and d:
            m = monthly[(d.year, d.month)]["call" if is_call else "spread"]
            m[0] += 1
            m[1] += r["amount"] or 0
    ws.auto_filter.ref = ws.dimensions

    # Итоги по месяцам: только подтверждённые оплаты, по месяцу даты записи
    ws2 = wb.create_sheet("Итоги по месяцам")
    _header(ws2, ["Месяц", "Расклады, шт", "Расклады, ₽", "Созвоны, шт", "Созвоны, ₽", "Всего, шт", "Всего, ₽"],
            [18, 13, 13, 13, 13, 11, 13])
    for (y, mth) in sorted(monthly, reverse=True):
        m = monthly[(y, mth)]
        ws2.append([f"{MONTHS[mth - 1]} {y}", m["spread"][0], m["spread"][1], m["call"][0], m["call"][1],
                    m["spread"][0] + m["call"][0], m["spread"][1] + m["call"][1]])
    last = ws2.max_row
    ws2.append(["Итого"] + [f"=SUM({get_column_letter(c)}2:{get_column_letter(c)}{last})" for c in range(2, 8)])
    for c in range(1, 8):
        ws2.cell(row=ws2.max_row, column=c).font = Font(bold=True)
    for row in ws2.iter_rows(min_row=2, min_col=3, max_col=7):
        for c in row:
            if c.column in (3, 5, 7):
                c.number_format = "# ##0"
    ws2.append([])
    ws2.append(["Считаются только подтверждённые записи, месяц — по дате записи."])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
